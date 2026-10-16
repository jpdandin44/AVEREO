"""Génère le Word du plan CIL depuis sa source Markdown, sans seconde rédaction.

Utiliser le Python du runtime documentaire Codex avec python-docx.
Le rendu visuel et la revue des pages restent requis après génération.
"""
from __future__ import annotations

import argparse
import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

APP = Path(__file__).resolve().parents[1]
TOKEN = re.compile(r'\[([^\]]+)\]\(([^)]+)\)|(\*\*[^*]+\*\*)|(`[^`]+`)')


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('**', '').replace('`', '')


def rich(paragraph, text, base, *, bold=False, size=None, color='000000'):
    position = 0
    for token in TOKEN.finditer(text):
        if token.start() > position:
            run = paragraph.add_run(text[position:token.start()])
            run.bold = bold
        if token.group(1):
            target = token.group(2)
            if not re.match(r'^https?://', target):
                target = (base / target).resolve().as_uri()
            rel = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
            run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            tint = OxmlElement('w:color'); tint.set(qn('w:val'), '1D4E63'); props.append(tint)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
            run.append(props); txt = OxmlElement('w:t'); txt.text = token.group(1); run.append(txt)
            link.append(run); paragraph._p.append(link)
        else:
            run = paragraph.add_run(token.group(0).strip('*`'))
            run.bold = bold or bool(token.group(3))
        position = token.end()
    if position < len(text):
        run = paragraph.add_run(text[position:]); run.bold = bold
    for run in paragraph.runs:
        run.font.name = 'Arial'; run.font.color.rgb = RGBColor.from_string(color)
        if size: run.font.size = Pt(size)


def blocks(content):
    lines = content.splitlines()
    if lines and lines[0] == '---':
        end = lines.index('---', 1); lines = lines[end+1:]
    cursor = 0
    while cursor < len(lines):
        line = lines[cursor].strip()
        if not line:
            cursor += 1; continue
        if line == '<!-- pagebreak -->':
            yield 'break', None; cursor += 1; continue
        if line.startswith('|'):
            rows = []
            while cursor < len(lines) and lines[cursor].strip().startswith('|'):
                row = [part.strip() for part in lines[cursor].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', part) for part in row): rows.append(row)
                cursor += 1
            yield 'table', rows; continue
        heading = re.match(r'^(#{1,3})\s+(.+)$', line)
        if heading:
            yield f'h{len(heading.group(1))}', heading.group(2); cursor += 1; continue
        bullet = re.match(r'^(- |\d+\. )(.+)$', line)
        prefix = bullet.group(1) if bullet else ''
        text = bullet.group(2) if bullet else line
        cursor += 1
        while cursor < len(lines) and lines[cursor].strip() and not re.match(r'^(#|\||<!--|[-] |\d+\. )', lines[cursor].strip()):
            text += ' ' + lines[cursor].strip(); cursor += 1
        yield 'bullet' if prefix == '- ' else 'number' if prefix else 'p', text


def field(paragraph, instruction):
    run = paragraph.add_run(); begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t'); value.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for item in [begin, text, separate, value, end]: run._r.append(item)


def table(document, rows, base):
    width = len(rows[0]); grid = document.add_table(rows=1, cols=width)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER; grid.autofit = False
    headers = rows[0]
    if headers[0] == 'Code': sizes = [1.6, 8.0, 1.9, 5.5]
    elif headers[0] == 'Essai': sizes = [1.35, 6.3, 9.35]
    elif headers[0] == 'Travail': sizes = [3.1, 6.4, 7.5]
    elif headers[0] == 'Domaine': sizes = [3.25, 5.35, 8.4]
    else: sizes = [3.4, 6.4, 7.2]
    for column, size in zip(grid.columns, sizes): column.width = Cm(size)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}'); border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
    grid._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = grid.rows[0] if index == 0 else grid.add_row()
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        for col, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(sizes[col]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, amount in [('top',90),('bottom',90),('left',105),('right',105)]:
                margin = OxmlElement(f'w:{side}'); margin.set(qn('w:w'), str(amount)); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
            props.append(margins)
            shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), '234B5A' if index == 0 else 'F2F5F6' if index % 2 == 0 else 'FFFFFF'); props.append(shading)
            paragraph = cell.paragraphs[0]; paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            paragraph.paragraph_format.keep_with_next = False
            if headers[0] in ['Code','Essai'] and col in ([0,2] if headers[0]=='Code' else [0]):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rich(paragraph, value, base, bold=index==0, size=9.5, color='FFFFFF' if index==0 else '000000')
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build(source, output):
    content = source.read_text(encoding='utf8')
    parsed = list(blocks(content)); document = Document()
    section = document.sections[0]; section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.85)
    section.header_distance = section.footer_distance = Cm(0.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Header','Footer','List Bullet','List Number']:
        style = document.styles[name]; style.font.name = 'Arial'; style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Arial')
        style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Arial')
        for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    normal = document.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08; normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [('Title',24,0,12),('Heading 1',15,7,9),('Heading 2',12,9,5)]:
        style = document.styles[name]; style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]; header.text = 'AVEREO     PASSEPORT IMMO'
    header.style = document.styles['Header']
    for run in header.runs: run.font.size = Pt(8.5)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Plan CIL  |  16 octobre 2026  |  Page '); field(footer,'PAGE'); footer.add_run(' sur '); field(footer,'NUMPAGES')
    for run in footer.runs: run.font.size = Pt(8.5)
    for kind, value in parsed:
        if kind == 'break': document.add_page_break()
        elif kind == 'table': table(document, value, source.parent)
        else:
            style = {'h1':'Title','h2':'Heading 1','h3':'Heading 2','bullet':'List Bullet','number':'List Number'}.get(kind)
            paragraph = document.add_paragraph(style=style)
            rich(paragraph, value, source.parent)
            if kind in ['bullet','number']: paragraph.paragraph_format.space_after = Pt(4)
    core = document.core_properties
    core.author = core.last_modified_by = 'AVEREO'
    core.title = 'Plan de développement du CIL dans Passeport Immo'
    core.subject = 'Évolutions proposées et poursuite de la version actuelle'
    core.comments = 'Source Markdown SHA256 ' + hashlib.sha256(content.encode('utf8')).hexdigest()
    core.created = core.modified = datetime(2026,10,16,tzinfo=timezone.utc)
    core.language = 'fr-FR'
    output.parent.mkdir(parents=True, exist_ok=True); document.save(output)
    actual = ' '.join(document.element.body.itertext())
    normalized = lambda text: re.sub(r'\s+', ' ', text).strip()
    actual = normalized(actual)
    for kind, value in parsed:
        if kind == 'break': continue
        values = [v for row in value for v in row] if kind == 'table' else [value]
        for item in values:
            if normalized(plain(item)) not in actual: raise ValueError('Contenu absent du Word : '+item[:100])
    print(f'Word généré et contenu vérifié : {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=APP/'docs/plan-action-cil.md')
    parser.add_argument('--output', type=Path, default=APP/'docs/livrables/plan-action-cil.docx')
    args = parser.parse_args(); build(args.source.resolve(), args.output.resolve())
